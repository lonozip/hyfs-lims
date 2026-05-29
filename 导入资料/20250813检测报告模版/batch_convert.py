import os
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree

BASE = r'd:\秦州_部门\孟令飞_LIMS\senaite.lims\senaite_lims\导入资料\20250813检测报告模版\20250813检测报告模版'
OUT_DIR = BASE

FILES = [
    '20250813含密封源仪表检测报告模版.docx',
    '20250813环境γ辐射剂量率检测报告模版.docx',
    '20250813医用X射线机房检测报告模版.docx',
    '20250813医用X射线性能检测报告模版.docx',
    '20250813职业性外照射个人检测报告模版.docx',
    '20250813职业性外照射年剂量检测评价报告模版.docx',
    '20250813X射线探伤机房检测报告模版.docx',
    '20250813X射线现场探伤检测报告模版.docx',
]

ALIGN_MAP = {
    WD_ALIGN_PARAGRAPH.CENTER: 'center',
    WD_ALIGN_PARAGRAPH.RIGHT: 'right',
    WD_ALIGN_PARAGRAPH.LEFT: 'left',
    WD_ALIGN_PARAGRAPH.JUSTIFY: 'justify',
    None: None
}

def get_alignment(p):
    return ALIGN_MAP.get(p.alignment, None)

def get_run_style_html(run):
    styles = []
    if run.bold:
        styles.append('font-weight:bold')
    if run.italic:
        styles.append('font-style:italic')
    if run.underline:
        styles.append('text-decoration:underline')
    if run.font.size:
        styles.append(f'font-size:{run.font.size.pt}pt')
    if run.font.name:
        styles.append(f'font-family:"{run.font.name}",SimSun,serif')
    if run.font.color and run.font.color.rgb:
        styles.append(f'color:#{run.font.color.rgb}')
    return '; '.join(styles)

def extract_images_from_element(xml_element, saved_images, img_rel_path):
    xml_str = etree.tostring(xml_element, encoding='unicode')
    imgs = []
    seen = set()
    for blip in re.finditer(r'r:embed="([^"]+)"', xml_str):
        embed = blip.group(1)
        if embed in saved_images and embed not in seen:
            seen.add(embed)
            imgs.append(f'<img src="{img_rel_path}/{saved_images[embed]}" style="max-width:100%;height:auto;display:block;" />')
    return imgs

def paragraph_to_html(p, saved_images, img_rel_path, skip_images=False):
    align = get_alignment(p)
    align_style = f'text-align:{align};' if align else ''
    imgs = [] if skip_images else extract_images_from_element(p._element, saved_images, img_rel_path)
    has_text = any(run.text.strip() for run in p.runs)

    if imgs and not has_text:
        return f'<div style="{align_style}">{"".join(imgs)}</div>'

    if not p.text.strip() and not imgs:
        return '<p style="min-height:1em;">&nbsp;</p>'

    runs_html = []
    for run in p.runs:
        text = run.text or ''
        text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;').replace('\n', '<br/>')
        s = get_run_style_html(run)
        runs_html.append(f'<span style="{s}">{text}</span>' if s else text)

    result = ''.join(runs_html)
    if imgs:
        result += ''.join(imgs)
    return f'<p style="{align_style}">{result}</p>'

def table_to_html(table, saved_images, img_rel_path):
    rows = table.rows
    if not rows:
        return ''

    cells_data = []
    for row in rows:
        row_cells = []
        for cell in row.cells:
            tc = cell._tc
            tc_pr = tc.find(qn('w:tcPr'))
            grid_span, vmerge = 1, None
            if tc_pr is not None:
                gs = tc_pr.find(qn('w:gridSpan'))
                if gs is not None:
                    grid_span = int(gs.get(qn('w:val')))
                vm = tc_pr.find(qn('w:vMerge'))
                if vm is not None:
                    vmerge = 'restart' if vm.get(qn('w:val')) == 'restart' else 'continue'

            cell_imgs = extract_images_from_element(tc, saved_images, img_rel_path)
            cell_html = ''.join(paragraph_to_html(p, saved_images, img_rel_path, skip_images=True) for p in cell.paragraphs)
            if cell_imgs:
                cell_html += '<div style="text-align:center;">' + ''.join(cell_imgs) + '</div>'
            row_cells.append({'html': cell_html, 'grid_span': grid_span, 'vmerge': vmerge})
        cells_data.append(row_cells)

    html = '<table class="report-table" style="border-collapse:collapse;width:100%;margin:8px 0;">'
    for ri, row_cells in enumerate(cells_data):
        html += '<tr>'
        for ci, cell in enumerate(row_cells):
            if cell['vmerge'] == 'continue':
                continue
            rowspan = 1
            if cell['vmerge'] == 'restart':
                rj = ri + 1
                while rj < len(cells_data) and ci < len(cells_data[rj]) and cells_data[rj][ci]['vmerge'] == 'continue':
                    rowspan += 1; rj += 1
            colspan = cell['grid_span'] if cell['grid_span'] > 1 else 1
            attrs = ''
            if colspan > 1: attrs += f' colspan="{colspan}"'
            if rowspan > 1: attrs += f' rowspan="{rowspan}"'
            html += f'<td style="border:1px solid #000;padding:4px 6px;vertical-align:middle;text-align:center;"{attrs}>{cell["html"]}</td>'
        html += '</tr>'
    html += '</table>'
    return html

def extract_body_elements(doc):
    body = doc.element.body
    elements = []
    pi, ti = 0, 0
    for child in body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p' and pi < len(doc.paragraphs):
            p = doc.paragraphs[pi]
            pPr = p._element.find(qn('w:pPr'))
            has_sect = pPr is not None and pPr.find(qn('w:sectPr')) is not None
            elements.append({'type': 'paragraph', 'data': p, 'section_break': has_sect})
            pi += 1
        elif tag == 'tbl' and ti < len(doc.tables):
            elements.append({'type': 'table', 'data': doc.tables[ti], 'section_break': False})
            ti += 1
    return elements

def smart_split_pages(elements):
    pages, current = [], []
    NEW_PAGE = ['四、检测结论', '附件', '报告说明', '一、仪器设备']

    def estimate_rows(els):
        r = 0
        for e in els:
            if e['type'] == 'table':
                r += len(e['data'].rows)
            else:
                t = e['data'].text.strip()
                if t: r += max(1, len(t)//60)
        return r

    def flush():
        nonlocal current
        if current:
            pages.append(current)
        current = []

    for el in elements:
        text = el['data'].text.strip() if el['type'] == 'paragraph' else ''

        if el['section_break']:
            if el['type'] == 'paragraph':
                current.append(el)
            flush()
            continue

        if any(text.startswith(h) for h in NEW_PAGE):
            flush()
            current = [el]
            continue

        if el['type'] == 'table':
            tr = len(el['data'].rows)
            ce = estimate_rows(current)
            if tr >= 10 and ce > 0:
                flush(); current = [el]; continue
            if ce + tr > 35 and ce > 0:
                flush(); current = [el]; continue
            tc = sum(1 for e in current if e['type'] == 'table')
            if tc >= 2 and tr > 3:
                flush(); current = [el]; continue

        current.append(el)

    flush()

    merged = []
    for i, page in enumerate(pages):
        has = any((e['type'] == 'paragraph' and e['data'].text.strip()) or e['type'] == 'table' for e in page)
        if not has and i+1 < len(pages):
            pages[i+1] = page + pages[i+1]
            continue
        merged.append(page)
    return merged

A4_CSS = """
        @page { size: A4; margin: 2cm 1.8cm 2cm 1.8cm; }
        @media print {
            body { margin: 0; padding: 0; }
            .page-container { width: 100%; max-width: 100%; margin: 0; padding: 0; box-shadow: none; page-break-after: always; }
            .page-container:last-child { page-break-after: auto; }
        }
        @media screen {
            body { background: #e8e8e8; display: flex; flex-direction: column; align-items: center; padding: 20px; }
            .page-container { width: 210mm; min-height: 297mm; background: #fff; padding: 2cm 1.8cm; margin-bottom: 30px; box-shadow: 0 2px 10px rgba(0,0,0,0.15); overflow-y: auto; }
        }
        * { box-sizing: border-box; }
        body { font-family: SimSun, "宋体", "Times New Roman", serif; font-size: 12pt; color: #000; line-height: 1.6; }
        .page-container p { margin: 0.3em 0; line-height: 1.6; }
        table.report-table { width: 100% !important; border-collapse: collapse; margin: 8px 0; font-size: 10.5pt; }
        table.report-table td, table.report-table th { border: 1px solid #000; padding: 4px 6px; vertical-align: middle; line-height: 1.5; }
        table.report-table p { margin: 2px 0; }
"""

def convert_docx(docx_path, out_path, img_dir_name):
    doc = Document(docx_path)
    img_dir = os.path.join(os.path.dirname(out_path), img_dir_name)
    os.makedirs(img_dir, exist_ok=True)

    saved_images = {}
    cnt = 0
    for rel_id, rel in doc.part.rels.items():
        if "image" in str(rel.reltype):
            cnt += 1
            ext = os.path.splitext(rel.target_ref)[1]
            fname = f'image_{cnt}{ext}'
            with open(os.path.join(img_dir, fname), 'wb') as f:
                f.write(rel.target_part.blob)
            saved_images[rel_id] = fname

    img_rel_path = img_dir_name
    elements = extract_body_elements(doc)
    pages = smart_split_pages(elements)

    base_name = os.path.splitext(os.path.basename(docx_path))[0]
    if base_name[:8].isdigit():
        report_title = base_name[8:]
    else:
        report_title = base_name

    combined = f'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{report_title}</title>
<style>
{A4_CSS}
</style>
</head>
<body>
'''

    for pi, page in enumerate(pages):
        body_html = ''
        for el in page:
            if el['type'] == 'paragraph':
                body_html += paragraph_to_html(el['data'], saved_images, img_rel_path)
            else:
                body_html += table_to_html(el['data'], saved_images, img_rel_path)
        combined += f'<!-- 第{pi+1}页 -->\n<div class="page-container">\n{body_html}\n</div>\n'

    combined += '</body>\n</html>'

    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(combined)

    return len(pages), cnt, len(combined)

for i, fname in enumerate(FILES):
    src = os.path.join(BASE, fname)
    name_no_ext = fname[8:-5] if fname[:8].isdigit() else fname[:-5]
    out_name = f'{name_no_ext}.html'
    out_path = os.path.join(OUT_DIR, out_name)
    img_dir_name = f'images_{name_no_ext}'

    print(f'[{i+1}/{len(FILES)}] {fname}')
    pages, imgs, size = convert_docx(src, out_path, img_dir_name)
    print(f'  -> {out_name} ({pages} pages, {imgs} images, {size} chars)')

print(f'\nDone!')
